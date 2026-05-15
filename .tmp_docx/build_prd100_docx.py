from __future__ import annotations

import math
import re
from pathlib import Path
from typing import List, Tuple

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path('/Users/jinmo/workspace/hitevision-digital-legion')
SRC = ROOT / 'docs/prd/PRD-100-数字军团OS总体产品需求文档-v1.0.md'
OUT = ROOT / 'docs/deliverables/PRD-100-数字军团OS总体产品需求文档-v1.0.docx'

ACCENT = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
SOFT = RGBColor(0xF7, 0xF9, 0xFC)
RULE = 'B7C3D0'
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_width(cell, dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.find(qn('w:tcW'))
    if tcw is None:
        tcw = OxmlElement('w:tcW')
        tc_pr.append(tcw)
    tcw.set(qn('w:w'), str(dxa))
    tcw.set(qn('w:type'), 'dxa')
    cell.width = Inches(dxa / 1440)


def set_table_geometry(table, widths: List[int], indent_dxa: int = TABLE_INDENT_DXA):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')

    tbl_ind = tbl_pr.find(qn('w:tblInd'))
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), str(indent_dxa))
    tbl_ind.set(qn('w:type'), 'dxa')

    tbl_layout = tbl_pr.find(qn('w:tblLayout'))
    if tbl_layout is None:
        tbl_layout = OxmlElement('w:tblLayout')
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn('w:type'), 'fixed')

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement('w:gridCol')
        grid_col.set(qn('w:w'), str(width))
        grid.append(grid_col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.find(qn('w:tcMar'))
            if tc_mar is None:
                tc_mar = OxmlElement('w:tcMar')
                tc_pr.append(tc_mar)
            for side, value in [('top', 80), ('bottom', 80), ('start', 120), ('end', 120)]:
                node = tc_mar.find(qn(f'w:{side}'))
                if node is None:
                    node = OxmlElement(f'w:{side}')
                    tc_mar.append(node)
                node.set(qn('w:w'), str(value))
                node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn('w:tblHeader'))
    if tbl_header is None:
        tbl_header = OxmlElement('w:tblHeader')
        tr_pr.append(tbl_header)
    tbl_header.set(qn('w:val'), 'true')


def set_run_font(run, name='Calibri', size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(7)
    pf.line_spacing = 1.15

    title = styles['Title']
    title.font.name = 'Calibri'
    title._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    title._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    title.font.size = Pt(23)
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.font.bold = True
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)

    subtitle = styles['Subtitle']
    subtitle.font.name = 'Calibri'
    subtitle._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    subtitle._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = MUTED
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)

    for style_name, size, color, before, after in [
        ('Heading 1', 18, RGBColor(0, 0, 0), 18, 10),
        ('Heading 2', 13, DARK, 12, 6),
        ('Heading 3', 11.5, ACCENT, 8, 4),
        ('Heading 4', 11, DARK, 6, 3),
    ]:
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]
        style.font.name = 'Calibri'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.12

    if 'Code Block' not in styles:
        code = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = 'Consolas'
        code._element.rPr.rFonts.set(qn('w:ascii'), 'Consolas')
        code._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
        code.font.size = Pt(9.5)
        code.paragraph_format.left_indent = Inches(0.25)
        code.paragraph_format.right_indent = Inches(0.15)
        code.paragraph_format.space_before = Pt(3)
        code.paragraph_format.space_after = Pt(3)
        code.paragraph_format.line_spacing = 1.0

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run('产品需求文档 | 数字军团OS')
    set_run_font(r, size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    rr = fp.add_run('第 ')
    set_run_font(rr, size=9, color=MUTED)
    add_page_field(fp)
    rr2 = fp.add_run(' 页')
    set_run_font(rr2, size=9, color=MUTED)


def add_heading_decor(paragraph, kind: str):
    p_pr = paragraph._p.get_or_add_pPr()
    pbdr = p_pr.find(qn('w:pBdr'))
    if pbdr is None:
        pbdr = OxmlElement('w:pBdr')
        p_pr.append(pbdr)
    for existing in list(pbdr):
        pbdr.remove(existing)
    if kind == 'h1':
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '8')
        bottom.set(qn('w:space'), '3')
        bottom.set(qn('w:color'), '2E74B5')
        pbdr.append(bottom)
    elif kind == 'h2':
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '8')
        left.set(qn('w:space'), '4')
        left.set(qn('w:color'), 'D0D7DE')
        pbdr.append(left)


def add_masthead(doc: Document):
    doc.add_paragraph('')
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run('PRD-100 数字军团OS总体产品需求文档')
    set_run_font(r, size=23, color=RGBColor(0, 0, 0), bold=True)

    p2 = doc.add_paragraph(style='Subtitle')
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r2 = p2.add_run('Phase 1 评审包总纲 | 正式评审稿')
    set_run_font(r2, size=14, color=MUTED)

    meta = [
        ('所属项目', '鸿合-数字军团OS'),
        ('当前版本', 'v1.0'),
        ('作者', 'Herick Liu'),
        ('创建日期', '2026-05-14'),
        ('文档定位', '总体总纲'),
        ('关联分册', 'PRD-101 / PRD-102 / PRD-103 / PRD-104 / PRD-105'),
    ]
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    table.autofit = False
    set_table_geometry(table, [4680, 4680], indent_dxa=0)
    flattened = meta[:]
    while len(flattened) < 6:
        flattened.append(('', ''))
    k = 0
    for row in table.rows:
        for cell_idx in range(2):
            cell = row.cells[cell_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            label, value = flattened[k]
            if label:
                a = p.add_run(label)
                set_run_font(a, size=9.5, color=MUTED, bold=True)
                a.add_break()
                b = p.add_run(value)
                set_run_font(b, size=11, color=INK, bold=False)
            k += 1
            set_cell_shading(cell, 'F7F9FC')
    doc.add_paragraph('')

    rule = doc.add_paragraph()
    p_pr = rule._p.get_or_add_pPr()
    pbdr = p_pr.find(qn('w:pBdr'))
    if pbdr is None:
        pbdr = OxmlElement('w:pBdr')
        p_pr.append(pbdr)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), RULE)
    pbdr.append(bottom)
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(12)
    rule.runs[0] if rule.runs else rule.add_run()
    rule.runs[-1].add_break(WD_BREAK.PAGE)


def extract_outline(lines: List[str]) -> List[Tuple[int, str]]:
    outline: List[Tuple[int, str]] = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith('#'):
            continue
        level = len(stripped) - len(stripped.lstrip('#'))
        text = stripped[level:].strip()
        outline.append((level, text))
    return outline


def add_toc(doc: Document, outline: List[Tuple[int, str]]):
    title = doc.add_paragraph(style='Heading 1')
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    parse_inline(title, '目录')
    add_heading_decor(title, 'h1')

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    parse_inline(intro, '以下目录用于评审阅读导航，按主章节和关键子章节展开。')

    for level, text in outline:
        if level == 1:
            continue
        if level > 3:
            continue
        p = doc.add_paragraph()
        indent = {2: 0.0, 3: 0.28}.get(level, 0.0)
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5 if level == 2 else 2.5)
        run = p.add_run(text)
        set_run_font(
            run,
            size=12.5 if level == 2 else 10.5,
            color=RGBColor(0, 0, 0) if level == 2 else DARK,
            bold=(level == 2),
        )
    end = doc.add_paragraph()
    end.add_run().add_break(WD_BREAK.PAGE)


def parse_inline(paragraph, text: str):
    text = text.replace('**', '__BOLD__', 0) if False else text
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name='Consolas', size=10, color=DARK)
        elif part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=11, color=INK, bold=True)
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=11, color=INK, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=11, color=INK)


def is_table_line(line: str) -> bool:
    return line.strip().startswith('|') and line.strip().endswith('|')


def split_table_row(line: str) -> List[str]:
    cells = [c.strip() for c in line.strip().strip('|').split('|')]
    return cells


def guess_widths(rows: List[List[str]]) -> List[int]:
    cols = max(len(r) for r in rows)
    weights = [1] * cols
    for c in range(cols):
        max_len = max((len(r[c]) if c < len(r) else 0) for r in rows)
        weights[c] = max(1, min(6, max_len))
    total = sum(weights)
    widths = [max(700, int(CONTENT_WIDTH_DXA * w / total)) for w in weights]
    diff = CONTENT_WIDTH_DXA - sum(widths)
    widths[-1] += diff
    return widths


def add_table(doc: Document, rows: List[List[str]]):
    widths = guess_widths(rows)
    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    table.autofit = False
    set_table_geometry(table, widths)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 10 else WD_ALIGN_PARAGRAPH.LEFT
            parse_inline(p, text)
            for run in p.runs:
                run.font.size = Pt(9.5 if i > 0 else 10)
                if i == 0:
                    run.bold = True
        if i == 0:
            set_repeat_table_header(table.rows[i])
            for cell in table.rows[i].cells:
                set_cell_shading(cell, 'F2F4F7')
    doc.add_paragraph('')


def add_code_block(doc: Document, lines: List[str]):
    for line in lines:
        p = doc.add_paragraph(style='Code Block')
        parse_inline(p, line)


def add_paragraph_with_style(doc: Document, text: str, style='Normal', quote=False):
    p = doc.add_paragraph(style=style)
    if quote:
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.1)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(8)
        p_pr = p._p.get_or_add_pPr()
        pbdr = p_pr.find(qn('w:pBdr'))
        if pbdr is None:
            pbdr = OxmlElement('w:pBdr')
            p_pr.append(pbdr)
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '10')
        left.set(qn('w:space'), '6')
        left.set(qn('w:color'), 'D0D7DE')
        pbdr.append(left)
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'F7F9FC')
        p_pr.append(shd)
    parse_inline(p, text)
    return p


def next_nonempty_line(lines: List[str], start: int) -> str | None:
    j = start
    while j < len(lines):
        if lines[j].strip():
            return lines[j].strip()
        j += 1
    return None


def add_list_item(doc: Document, text: str, ordered=False, level=0):
    p = doc.add_paragraph(style='List Number' if ordered else 'List Bullet')
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    parse_inline(p, text)


def build():
    doc = Document()
    configure_document(doc)
    add_masthead(doc)

    lines = SRC.read_text(encoding='utf-8').splitlines()
    outline = extract_outline(lines)
    add_toc(doc, outline)
    i = 0
    in_code = False
    code_lines: List[str] = []
    skipped_title = False
    seen_major_heading = False
    while i < len(lines):
        line = lines[i]
        stripped = line.rstrip('\n')
        if stripped.startswith('```'):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(stripped)
            i += 1
            continue
        if not stripped.strip():
            i += 1
            continue
        if stripped.strip() in ('***', '---'):
            upcoming = next_nonempty_line(lines, i + 1)
            if upcoming and upcoming.startswith('## '):
                i += 1
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_pr = p._p.get_or_add_pPr()
            pbdr = p_pr.find(qn('w:pBdr'))
            if pbdr is None:
                pbdr = OxmlElement('w:pBdr')
                p_pr.append(pbdr)
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'DADCE0')
            pbdr.append(bottom)
            i += 1
            continue
        if is_table_line(stripped):
            rows = []
            while i < len(lines) and is_table_line(lines[i]):
                row = split_table_row(lines[i])
                if not all(re.fullmatch(r':?-{3,}:?', c.replace(' ', '')) for c in row):
                    rows.append(row)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            text = stripped[level:].strip()
            if level == 1 and not skipped_title:
                skipped_title = True
                i += 1
                continue
            style = {2: 'Heading 1', 3: 'Heading 2', 4: 'Heading 3', 5: 'Heading 4'}.get(level, 'Heading 4')
            p = doc.add_paragraph(style=style)
            if level == 2 and seen_major_heading:
                p.paragraph_format.page_break_before = True
            parse_inline(p, text)
            if level == 2:
                seen_major_heading = True
                add_heading_decor(p, 'h1')
            elif level == 3:
                add_heading_decor(p, 'h2')
            i += 1
            continue
        if stripped.lstrip().startswith('>'):
            quote_text = stripped.lstrip()[1:].strip()
            add_paragraph_with_style(doc, quote_text, quote=True)
            i += 1
            continue
        m_num = re.match(r'^(\s*)(\d+)\.\s+(.*)$', stripped)
        if m_num:
            level = max(0, len(m_num.group(1)) // 2)
            add_list_item(doc, m_num.group(3), ordered=True, level=level)
            i += 1
            continue
        m_bullet = re.match(r'^(\s*)[-*]\s+(.*)$', stripped)
        if m_bullet:
            level = max(0, len(m_bullet.group(1)) // 2)
            add_list_item(doc, m_bullet.group(2), ordered=False, level=level)
            i += 1
            continue
        # merge consecutive plain lines into a single paragraph
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if not nxt.strip():
                i += 1
                break
            if nxt.startswith('#') or nxt.strip() in ('***', '---') or nxt.strip().startswith('>') or is_table_line(nxt) or re.match(r'^(\s*)(\d+)\.\s+', nxt) or re.match(r'^(\s*)[-*]\s+', nxt) or nxt.startswith('```'):
                break
            para_lines.append(nxt.strip())
            i += 1
        add_paragraph_with_style(doc, ' '.join(para_lines))

    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
